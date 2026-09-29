import streamlit as st
import plotly.graph_objects as go
import plotly.express as px
import pandas as pd
import graphviz
import math
import google.generativeai as genai
from PIL import Image
import os
import zipfile
import base64

# Import from our custom files
from utils import (
    extract_text_from_file, load_knowledge_base, convert_df_to_csv, 
    create_pdf, create_pptx, load_database, save_to_database, scrape_url, KB_DIR
)
from evori import apply_theme_and_cat

from pptx import Presentation 

st.set_page_config(page_title="AI Academic Visualizer", layout="wide")

# ==========================================
# DIRECT DOCUMENT LINKING (Listens to Evori)
# ==========================================
try:
    target_doc = st.query_params.get("doc", None)
except AttributeError:
    params = st.experimental_get_query_params()
    target_doc = params.get("doc", [None])[0]

def read_pptx_slides(file_path):
    try:
        prs = Presentation(file_path)
        slides = []
        for slide in prs.slides:
            content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    content.append(shape.text.strip())
            slides.append("\n\n".join(content))
        return slides
    except Exception:
        return []

# ==========================================
# PRE-LOAD FILES FOR EVORI'S BRAIN
# ==========================================
global_valid_files = []
if os.path.exists(KB_DIR):
    for root, dirs, files in os.walk(KB_DIR):
        for f in files:
            if not f.startswith('.'):
                global_valid_files.append(os.path.relpath(os.path.join(root, f), KB_DIR))

# ==========================================
# UI SIDEBAR
# ==========================================
st.sidebar.title("🎨 Appearance")
app_theme = st.sidebar.selectbox("Dashboard Theme", [
    "Dark Mode (Default)", "Light Mode", "Evori Neon (Pink)",
    "Evori Neon (Purple)", "Evori Neon (Green)", "Evori Neon (Orange)"
])
st.sidebar.markdown("---")

st.sidebar.title("Toolbox Navigation")
st.sidebar.markdown("---")

cat_equipped = st.sidebar.toggle("🐱 Equip Evori Familiar", value=True)
st.sidebar.markdown("---")

if "api_key" not in st.session_state: st.session_state.api_key = ""
temp_api_key = st.sidebar.text_input("Enter your Google Gemini API Key:", type="password", value=st.session_state.api_key)
if st.sidebar.button("Save API Key"):
    st.session_state.api_key = temp_api_key
    st.toast("API Key Locked In securely!", icon="🔑")

st.sidebar.markdown("---")

st.sidebar.title("📚 Academic Reference Library")
zip_upload = st.sidebar.file_uploader("Upload Multi-Format .zip", type=["zip"])
if zip_upload is not None:
    try:
        with zipfile.ZipFile(zip_upload, "r") as z: z.extractall(KB_DIR)
        st.cache_data.clear() 
        st.sidebar.success("✅ Multi-Format Library Updated!")
        st.toast("Theory library synchronized!", icon="📚")
        st.rerun() 
    except Exception as e: st.sidebar.error(f"Error extracting zip: {e}")

kb_text_content, kb_image_list = load_knowledge_base()

if kb_text_content or kb_image_list:
    st.sidebar.caption(f"🟢 Active: {len(kb_text_content)} chars | {len(kb_image_list)} images")
else: st.sidebar.caption("⚪ Library: Empty")

st.sidebar.markdown("---")

cat_options = ["Strategy & Marketing", "Finance & Supply Chain", "Project Management", "Theories Library"]

# If Evori clicked a link, automatically open the Theories Library!
default_cat_idx = 3 if target_doc else 0
category = st.sidebar.selectbox("Select a Module", cat_options, index=default_cat_idx)

# Clear the URL memory if the user navigates away manually
if category != "Theories Library":
    try:
        if "doc" in st.query_params:
            del st.query_params["doc"]
    except AttributeError:
        st.experimental_set_query_params()

if category == "Strategy & Marketing": framework = st.sidebar.radio("Available Frameworks:", ["AI Analysis Workspace", "BCG Portfolio Matrix", "Board of Directors Simulation"])
elif category == "Finance & Supply Chain": framework = st.sidebar.radio("Available Frameworks:", ["Financial Break-Even", "Financial Health Radar", "NPV & Cash Flow Projection", "Project Budget Planner", "EOQ Inventory Optimization", "Kraljic Portfolio Matrix", "JIT Supply Chain (Diagram)"])
elif category == "Project Management": framework = st.sidebar.radio("Available Frameworks:", ["Action Plan (Gantt Chart)", "Risk Assessment Heatmap", "Stakeholder Power/Interest Grid"])
else: framework = "Theories Viewer"

st.sidebar.markdown("---")
st.sidebar.title("📁 My Projects Library")
db = load_database()
if db:
    saved_projects = list(db.keys())
    selected_project = st.sidebar.selectbox("Load a saved project:", ["-- Select --"] + saved_projects)
    if st.sidebar.button("Load Project"):
        if selected_project != "-- Select --":
            st.session_state.task_input = selected_project; st.session_state.ai_report = db[selected_project]; st.session_state.qa_history = []
            st.toast(f"Loaded '{selected_project}'!", icon="📂"); st.rerun()

# ==========================================
# CAT INJECTION & THEMING
# ==========================================
current_cat_mode = "Analyst"
if framework == "Board of Directors Simulation": current_cat_mode = "CEO"
elif category == "Finance & Supply Chain": current_cat_mode = "Finance"
elif category == "Project Management": current_cat_mode = "Project"
elif framework == "BCG Portfolio Matrix": current_cat_mode = "Strategist"
elif framework == "Theories Viewer": current_cat_mode = "Examiner"

apply_theme_and_cat(app_theme, current_cat_mode, cat_equipped, global_valid_files)

# ==========================================
# STATE MEMORY
# ==========================================
if "ai_report" not in st.session_state: st.session_state.ai_report = ""
if "task_input" not in st.session_state: st.session_state.task_input = ""
if "audience" not in st.session_state: st.session_state.audience = "University Thesis"
if "competitor_url" not in st.session_state: st.session_state.competitor_url = ""
if "qa_history" not in st.session_state: st.session_state.qa_history = []

# ==========================================
# VIEWER
# ==========================================
if category == "Theories Library":
    st.title("📚 Master Theory & Reference Library")
    st.write("Browse uploaded reference books, notes, and visual frameworks.")
    st.markdown("---")
    
    if global_valid_files:
        sorted_files = sorted(global_valid_files, key=str.casefold)
        
        # If Evori clicked a link, pre-select that document
        doc_idx = 0
        if target_doc and target_doc in sorted_files:
            doc_idx = sorted_files.index(target_doc)
        
        selected_theory = st.selectbox("🔍 Search and Select a Document (A-Z):", sorted_files, index=doc_idx)
        
        # Keep URL matching the currently selected document
        try:
            if selected_theory:
                st.query_params["doc"] = selected_theory
        except AttributeError:
            st.experimental_set_query_params(doc=selected_theory)
        
        if selected_theory:
            file_path = os.path.join(KB_DIR, selected_theory)
            st.markdown(f"### 📄 `{selected_theory}`")
            
            if selected_theory.lower().endswith(('.png', '.jpg', '.jpeg')):
                try: st.image(Image.open(file_path), caption=selected_theory, use_container_width=True)
                except Exception as e: st.error(f"Image error: {e}")
            
            elif selected_theory.lower().endswith('.pdf'):
                with open(file_path, "rb") as f:
                    base64_pdf = base64.b64encode(f.read()).decode('utf-8')
                pdf_display = f'<iframe src="data:application/pdf;base64,{base64_pdf}" width="100%" height="800" type="application/pdf"></iframe>'
                st.markdown(pdf_display, unsafe_allow_html=True)
            
            elif selected_theory.lower().endswith(('.xls', '.xlsx')):
                st.markdown("### 📊 Interactive Excel Viewer")
                try:
                    excel_file = pd.ExcelFile(file_path)
                    sheet_names = excel_file.sheet_names
                    if len(sheet_names) > 1:
                        selected_sheet = st.selectbox("Select Excel Sheet:", sheet_names)
                    else:
                        selected_sheet = sheet_names[0]
                    
                    df = pd.read_excel(file_path, sheet_name=selected_sheet)
                    st.dataframe(df, use_container_width=True)
                except Exception as e:
                    st.error(f"Error reading Excel file: {e}")

            elif selected_theory.lower().endswith(('.ppt', '.pptx')):
                st.warning("⚠️ **Browser Limitation:** Web browsers physically cannot render the visual graphics of a PowerPoint file. To see the actual visual slides, **Save As -> PDF** in PowerPoint and upload that PDF here!")
                slides = read_pptx_slides(file_path)
                if slides:
                    st.markdown("### 📝 Text-Only Slide Extraction")
                    slide_idx = st.slider("Navigate Slides", 1, len(slides), 1) - 1
                    
                    st.markdown(f"**Slide {slide_idx + 1} of {len(slides)}**")
                    st.info(slides[slide_idx] if slides[slide_idx] else "[No text on this slide]")
                    
                    st.markdown("<br>", unsafe_allow_html=True)
                    with open(file_path, "rb") as f:
                        st.download_button("📥 Download Original .pptx", f, file_name=selected_theory)
                else:
                    st.warning("Could not extract slides or the presentation contains no text.")
                    
            else:
                preview = extract_text_from_file(file_path)
                if preview: st.text_area("Preview", preview, height=500)
                else: st.warning("No readable text found.")
    else: 
        st.info("Upload a `.zip` using the sidebar!")

else:
    st.title("Interactive AI Academic Dashboard")
    col_task, col_tone = st.columns([2, 1])
    with col_task: temp_task = st.text_input("Project Topic:", value=st.session_state.task_input, placeholder="e.g., A premium electric scooter sharing service...")
    with col_tone: temp_audience = st.selectbox("Target Audience & Tone:", ["University Thesis", "Venture Capital Pitch", "Internal Board Memo"], index=["University Thesis", "Venture Capital Pitch", "Internal Board Memo"].index(st.session_state.audience))

    st.markdown("### 📎 Context & References")
    col_url, col_file = st.columns(2)
    with col_url: temp_url = st.text_input("Live Competitor URL (Scrape website data):", value=st.session_state.competitor_url)
    with col_file: uploaded_file = st.file_uploader("Upload session reference (.txt, .pdf, .png)", type=["txt", "png", "jpg", "jpeg", "pdf", "docx"])

    col_btn1, col_btn2 = st.columns(2)
    with col_btn1:
        if st.button("Submit Prompt", use_container_width=True):
            st.session_state.task_input = temp_task; st.session_state.audience = temp_audience; st.session_state.competitor_url = temp_url; st.session_state.qa_history = [] 
            st.toast("Prompt locked!", icon="🎯")
            if st.session_state.cat_equipped: components.html("<script>if(window.parent.evoriActions) window.parent.evoriActions.speak('Action', 'Prompt saved! 👍');</script>", height=0)
    with col_btn2:
        if st.button("Submit & Save Project", use_container_width=True):
            st.session_state.task_input = temp_task; st.session_state.audience = temp_audience; st.session_state.competitor_url = temp_url; st.session_state.qa_history = []
            save_to_database(temp_task, st.session_state.ai_report)
            st.toast("Prompt saved to Library!", icon="💾")
            if st.session_state.cat_equipped: components.html("<script>if(window.parent.evoriActions) window.parent.evoriActions.speak('Action', 'Saved to disk! 💾');</script>", height=0)

    st.markdown("---")

    if st.session_state.task_input: st.subheader(f"Current Project: {st.session_state.task_input}")
    else: st.subheader(f"Viewing: {framework}")

    def generate_report(raw_topic, analysis_context):
        if not st.session_state.api_key: st.error("⚠️ Please enter your API Key in the sidebar!"); return
        if not st.session_state.task_input: st.error("⚠️ Please enter a project topic and click 'Submit Prompt'!"); return
        try:
            genai.configure(api_key=st.session_state.api_key)
            model = genai.GenerativeModel('gemini-1.5-flash')
            with st.spinner(f"🤠 Prompt Cowboy is analyzing..."):
                tone_guide = {
                    "University Thesis": "Rigorously academic, theoretical, highly structured. Benchmark findings against the academic theories and literature provided in the reference library.",
                    "Venture Capital Pitch": "Punchy, ROI-focused. Highlight scalability, moat, and financial upside.",
                    "Internal Board Memo": "Concise, operational, risk-focused executive summary."
                }
                ai_input = [f"You are 'Prompt Cowboy'. Write a {analysis_context} for: '{raw_topic}'. TONE: {st.session_state.audience}. STYLE: {tone_guide.get(st.session_state.audience, '')} BENCHMARK using library."]
                if kb_text_content: ai_input.append(f"\n\n[ACADEMIC BENCHMARKING LIBRARY (TEXT)]:\n{kb_text_content}")
                if kb_image_list:
                    ai_input.append("\n\n[VISUAL FRAMEWORKS]:")
                    for img in kb_image_list:
                        try: ai_input.append(Image.open(img))
                        except Exception: pass
                if st.session_state.competitor_url:
                    scraped = scrape_url(st.session_state.competitor_url)
                    if not scraped.startswith("[Error"): ai_input.append(f"\n\nCompetitor Data:\n{scraped}")
                if uploaded_file is not None:
                    if uploaded_file.name.lower().endswith(('.png', '.jpg', '.jpeg')): ai_input.append(Image.open(uploaded_file))
                    else: ai_input.append(f"\n\nSession Doc:\n{extract_text_from_file(uploaded_file.name)}")

                final_response = model.generate_content(ai_input)
                st.session_state.ai_report = final_response.text
                st.session_state.qa_history = [] 
            st.toast("Analysis Complete!", icon="🚀")
            if st.session_state.cat_equipped: 
                components.html("<script>if(window.parent.evoriActions) window.parent.evoriActions.speak('Action', 'Analysis Complete! ✨');</script>", height=0)
        except Exception as e: st.error(f"Error: {e}")

    if framework == "AI Analysis Workspace":
        if st.button("Generate Complete Strategic Master Report"): generate_report(st.session_state.task_input, "comprehensive ALL-IN-ONE Strategic Master Report containing a full SWOT, PESTLE, STP, Porter's Five Forces, and VRIO Framework")
        analysis_type = st.selectbox("Which individual analysis do you need?", ["SWOT Analysis", "PESTLE Analysis", "Porter's Five Forces", "STP Analysis", "VRIO Framework"])
        if st.button(f"Generate Single: {analysis_type}"): generate_report(st.session_state.task_input, analysis_type)

    elif framework == "BCG Portfolio Matrix":
        col1, col2, col3 = st.columns(3)
        with col1: p1 = st.text_input("P1", "Product A"); s1 = st.slider("P1 Share", 0, 100, 80); g1 = st.slider("P1 Growth", 0, 100, 80)
        with col2: p2 = st.text_input("P2", "Product B"); s2 = st.slider("P2 Share", 0, 100, 20); g2 = st.slider("P2 Growth", 0, 100, 80)
        with col3: p3 = st.text_input("P3", "Product C"); s3 = st.slider("P3 Share", 0, 100, 80); g3 = st.slider("P3 Growth", 0, 100, 20)
        bcg_df = pd.DataFrame({"Product": [p1, p2, p3], "Share": [s1, s2, s3], "Growth": [g1, g2, g3], "Size": [40, 40, 40]})
        fig = px.scatter(bcg_df, x="Share", y="Growth", text="Product", size="Size", color="Product", range_x=[0, 100], range_y=[0, 100])
        fig.add_hline(y=50, line_dash="dash", line_color="rgba(255,255,255,0.2)"); fig.add_vline(x=50, line_dash="dash", line_color="rgba(255,255,255,0.2)")
        fig.update_layout(xaxis_title="Relative Market Share", yaxis_title="Market Growth Rate", showlegend=False, paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA'))
        st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate AI BCG Strategy"): generate_report(st.session_state.task_input, f"BCG Matrix analysis for: {p1}({s1}%,{g1}%), {p2}({s2}%,{g2}%), {p3}({s3}%,{g3}%).")

    elif framework == "Board of Directors Simulation":
        st.write("### 1. The Executive Boardroom")
        if st.button("Consult the Board"): generate_report(st.session_state.task_input, "multi-agent simulation where three personas aggressively debate the project: a Skeptical CFO, an Aggressive CMO, and a Pragmatic COO. Present as a lively transcript.")

    elif framework == "Financial Break-Even":
        c1, c2, c3 = st.columns(3); fc = c1.slider("Fixed Costs (€)", 10000, 500000, 280000); sp = c2.slider("Selling Price (€)", 10, 500, 150); vc = c3.slider("Variable Cost (€)", 5, 400, 60)
        units = list(range(0, 5000, 250)); rev = [sp * u for u in units]; cost = [fc + (vc * u) for u in units]
        try: be = fc / (sp - vc)
        except ZeroDivisionError: be = 0
        fig = go.Figure(); fig.add_trace(go.Scatter(x=units, y=rev, name='Revenue', line=dict(color='#00FFAA', width=3))); fig.add_trace(go.Scatter(x=units, y=cost, name='Costs', line=dict(color='#FF4B4B', width=3)))
        if 0 < be <= max(units): fig.add_vline(x=be, line_dash="dash", line_color="white", annotation_text=f" BE: {int(be)}")
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Break-Even Report"): generate_report(st.session_state.task_input, "financial break-even analysis")

    elif framework == "Financial Health Radar":
        c1, c2, c3, c4, c5 = st.columns(5); liq = c1.slider("Liquidity", 1, 10, 7); prof = c2.slider("Profitability", 1, 10, 6); mkt = c3.slider("Market Share", 1, 10, 5); brand = c4.slider("Brand Loyalty", 1, 10, 8); ops = c5.slider("Ops Efficiency", 1, 10, 7)
        fig = go.Figure(); fig.add_trace(go.Scatterpolar(r=[liq, prof, mkt, brand, ops], theta=['Liquidity', 'Profitability', 'Market Share', 'Brand Loyalty', 'Ops Efficiency'], fill='toself', name='Health', line_color='#00FFAA'))
        fig.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 10], color='white'), bgcolor='rgba(0,0,0,0)'), showlegend=False, paper_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate AI Radar Analysis"): generate_report(st.session_state.task_input, f"strategic health radar analysis. Scores: Liq({liq}), Prof({prof}), Mkt({mkt}), Brand({brand}), Ops({ops}).")

    elif framework == "NPV & Cash Flow Projection":
        c1, c2, c3, c4 = st.columns(4); inv = c1.number_input("Investment (€)", value=150000); ret = c2.number_input("Year 1 Return (€)", value=35000); gr = c3.slider("Growth (%)", -10, 100, 15); dr = c4.slider("Discount (%)", 1, 20, 8)
        yrs = ["Y0", "Y1", "Y2", "Y3", "Y4", "Y5"]; cf = [-inv]; cum = [-inv]; npv = -inv
        for t in range(1, 6): c = ret * ((1 + (gr / 100.0)) ** (t - 1)); cf.append(c); cum.append(cum[-1] + c); npv += c / ((1 + (dr / 100.0)) ** t)
        fig = go.Figure(); fig.add_trace(go.Bar(x=yrs, y=cf, name='Cash Flow', marker_color=['#FF4B4B' if x < 0 else '#00FFAA' for x in cf])); fig.add_trace(go.Scatter(x=yrs, y=cum, name='Cumulative', line=dict(color='orange', width=4)))
        fig.update_layout(title=f"<b>NPV: €{npv:,.2f}</b>", paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate AI NPV Report"): generate_report(st.session_state.task_input, f"NPV analysis. Inv: €{inv}, Y1: €{ret} (growth {gr}%), discount: {dr}%. NPV: €{npv:,.2f}.")

    elif framework == "Project Budget Planner":
        c1, c2 = st.columns(2)
        with c1: p = st.number_input("Personnel (€)", value=50000); e = st.number_input("Equipment (€)", value=25000); m = st.number_input("Marketing (€)", value=15000)
        with c2: o = st.number_input("Operations (€)", value=20000); l = st.number_input("Legal (€)", value=5000); c = st.number_input("Contingency (€)", value=10000)
        df = pd.DataFrame({"Category": ["Personnel", "Equipment", "Marketing", "Operations", "Legal", "Contingency"], "Amount": [p, e, m, o, l, c]})
        fig = px.pie(df[df["Amount"] > 0], values='Amount', names='Category', hole=0.4, color_discrete_sequence=px.colors.qualitative.Set3)
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Budget Report"): generate_report(st.session_state.task_input, f"budget justification for €{p+e+m+o+l+c} total.")

    elif framework == "EOQ Inventory Optimization":
        c1, c2, c3 = st.columns(3); d = c1.slider("Demand", 1000, 100000, 15000); o = c2.slider("Order Cost", 10, 1000, 200); h = c3.slider("Hold Cost", 1.0, 50.0, 5.0)
        eoq = math.sqrt((2 * d * o) / h)
        q = list(range(100, int(eoq * 3), 50)); hc = [(x / 2) * h for x in q]; oc = [(d / x) * o for x in q]; tc = [x + y for x, y in zip(hc, oc)]
        fig = go.Figure(); fig.add_trace(go.Scatter(x=q, y=hc, name='Hold Cost', line=dict(color='orange'))); fig.add_trace(go.Scatter(x=q, y=oc, name='Order Cost', line=dict(color='#00FFAA'))); fig.add_trace(go.Scatter(x=q, y=tc, name='Total Cost', line=dict(color='#FF4B4B')))
        fig.add_vline(x=eoq, line_color="white", annotation_text=f" EOQ: {int(eoq)}")
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Inventory Report"): generate_report(st.session_state.task_input, "inventory management EOQ strategy.")

    elif framework == "Kraljic Portfolio Matrix":
        fig = px.scatter(x=[25, 75, 25, 75], y=[75, 75, 25, 25], text=["Leverage", "Strategic", "Non-Critical", "Bottleneck"], size=[40, 40, 40, 40], color=["#00FFAA", "#FF4B4B", "gray", "orange"])
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Kraljic Analysis"): generate_report(st.session_state.task_input, "Kraljic Portfolio procurement strategy.")

    elif framework == "JIT Supply Chain (Diagram)":
        g = graphviz.Digraph(); g.attr(bgcolor='transparent', fontcolor='white'); g.node('A', 'Raw Materials', fontcolor='white'); g.node('B', 'Manufacturing', fontcolor='white'); g.edge('A', 'B', fontcolor='white'); st.graphviz_chart(g)
        if st.button("Generate JIT Report"): generate_report(st.session_state.task_input, "JIT logistics strategy.")

    elif framework == "Action Plan (Gantt Chart)":
        df = pd.DataFrame([dict(Task="Initiation", Start='2026-10-01', Finish='2026-11-15'), dict(Task="Procurement", Start='2026-11-10', Finish='2027-01-20')])
        fig = px.timeline(df, x_start="Start", x_end="Finish", y="Task", color="Task")
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Project Plan"): generate_report(st.session_state.task_input, "project action plan timeline.")

    elif framework == "Risk Assessment Heatmap":
        fig = go.Figure(data=go.Heatmap(z=[[1,2],[3,4]], colorscale='RdYlGn_r'))
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Risk Register"): generate_report(st.session_state.task_input, "project risk register and mitigation.")

    elif framework == "Stakeholder Power/Interest Grid":
        df = pd.DataFrame({"Stakeholder": ["Gov", "Public"], "Interest": [60, 85], "Power": [90, 40], "Size": [40, 40]})
        fig = px.scatter(df, x="Interest", y="Power", text="Stakeholder", size="Size", color="Stakeholder", range_x=[0, 100], range_y=[0, 100])
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font=dict(color='#FAFAFA')); st.plotly_chart(fig, use_container_width=True)
        if st.button("Generate Stakeholder Plan"): generate_report(st.session_state.task_input, "stakeholder management plan.")

    if st.session_state.ai_report:
        st.markdown("---")
        export_content = f"PROJECT TOPIC: {st.session_state.task_input}\nTONE: {st.session_state.audience}\n" + "="*50 + "\n\n" + st.session_state.ai_report
        colA, colB, colC, colD = st.columns(4)
        with colA: st.download_button("📝 Text (.txt)", data=export_content, file_name="Report.txt", use_container_width=True)
        with colB: st.download_button("📄 PDF (.pdf)", data=create_pdf(st.session_state.task_input, st.session_state.ai_report), file_name="Report.pdf", use_container_width=True)
        with colC: st.download_button("📊 PowerPoint (.pptx)", data=create_pptx(st.session_state.task_input, st.session_state.ai_report), file_name="AI_Presentation.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", use_container_width=True)
        with colD:
            if st.button("💾 Save to Library", use_container_width=True): save_to_database(st.session_state.task_input, st.session_state.ai_report); st.toast("Saved to sidebar library!", icon="💾")
        st.markdown("## 📄 Your AI-Generated Report")
        st.markdown(st.session_state.ai_report)
        st.markdown("---")
        st.markdown("## 🎤 Q&A Defense Simulator")
        for chat in st.session_state.qa_history:
            if chat["role"] == "user": st.markdown(f"**👤 You:** {chat['text']}")
            else: st.markdown(f"**🤖 AI Reviewer:** {chat['text']}")
        qa_input = st.text_input("Type your question (e.g., 'Grill me on my budget'):")
        if st.button("Ask / Defend"):
            if qa_input and st.session_state.api_key:
                st.session_state.qa_history.append({"role": "user", "text": qa_input})
                with st.spinner("Formulating response..."):
                    try:
                        genai.configure(api_key=st.session_state.api_key)
                        qa_model = genai.GenerativeModel('gemini-1.5-flash')
                        defense_prompt = f"Act as a skeptical {st.session_state.audience}. Report: {st.session_state.ai_report}. User asks: '{qa_input}'. Challenge their logic."
                        st.session_state.qa_history.append({"role": "ai", "text": qa_model.generate_content([defense_prompt]).text})
                        st.rerun() 
                    except Exception as e: st.error(f"Error: {e}")