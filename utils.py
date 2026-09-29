import os
import json
import requests
from bs4 import BeautifulSoup
from io import BytesIO
from fpdf import FPDF
from pptx import Presentation
import PyPDF2
import docx
import pandas as pd
import streamlit as st

KB_DIR = "knowledge_base"
DB_FILE = "projects.json"

os.makedirs(KB_DIR, exist_ok=True)

def extract_text_from_file(file_path):
    text = ""
    file_lower = file_path.lower()
    try:
        if file_lower.endswith(('.txt', '.md', '.csv', '.json')):
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        elif file_lower.endswith('.pdf'):
            reader = PyPDF2.PdfReader(file_path)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        elif file_lower.endswith(('.doc', '.docx')):
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif file_lower.endswith(('.ppt', '.pptx')):
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif file_lower.endswith(('.xls', '.xlsx')):
            df = pd.read_excel(file_path)
            text = df.to_string()
    except Exception as e:
        text = f"[Error extracting text: {e}]"
    return text

# NEW: Reads PPTX slide by slide for the presentation viewer
def read_pptx_slides(file_path):
    try:
        prs = Presentation(file_path)
        slides = []
        for slide in prs.slides:
            content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    content.append(shape.text.strip())
            slides.append("\n\n".join(content))
        return slides
    except Exception:
        return []

@st.cache_data(show_spinner=False)
def load_knowledge_base():
    kb_text = ""
    kb_images = []
    if os.path.exists(KB_DIR):
        for root, dirs, files in os.walk(KB_DIR):
            for file in files:
                file_path = os.path.join(root, file)
                file_lower = file.lower()
                if file_lower.endswith(('.jpg', '.jpeg', '.png')):
                    kb_images.append(file_path)
                elif file_lower.endswith(('.txt', '.md', '.csv', '.json', '.pdf', '.doc', '.docx', '.ppt', '.pptx', '.xls', '.xlsx')):
                    extracted = extract_text_from_file(file_path)
                    if extracted.strip():
                        kb_text += f"\n\n=== [REFERENCE LIBRARY SOURCE: {file}] ===\n" + extracted
    return kb_text, kb_images

@st.cache_data
def convert_df_to_csv(df): 
    return df.to_csv(index=False).encode('utf-8')

def create_pdf(title, content):
    pdf = FPDF(); pdf.add_page(); pdf.set_font("Arial", size=12)
    pdf.set_font("Arial", 'B', 16); safe_title = title.encode('latin-1', 'replace').decode('latin-1')
    pdf.cell(0, 10, txt=f"Project Topic: {safe_title}", ln=True); pdf.ln(5)
    pdf.set_font("Arial", size=11); safe_content = content.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 7, txt=safe_content)
    return bytes(pdf.output(dest='S'), encoding='latin-1')

def create_pptx(title, content):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title; slide.placeholders[1].text = "Automated AI Strategy Presentation"
    for para in content.split('\n\n'):
        if len(para.strip()) > 15:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "Key Insight"; s.placeholders[1].text_frame.text = para.replace('**', '').replace('*', '').strip()
    stream = BytesIO(); prs.save(stream); stream.seek(0)
    return stream.getvalue()

def load_database():
    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r", encoding="utf-8") as f: return json.load(f)
    return {}

def save_to_database(project_name, report_text):
    db = load_database(); db[project_name] = report_text
    with open(DB_FILE, "w", encoding="utf-8") as f: json.dump(db, f, indent=4)

def scrape_url(url):
    try:
        response = requests.get(url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=10)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        return " ".join([elem.get_text(strip=True) for elem in soup.find_all(['p', 'h1', 'h2', 'h3', 'li'])])[:10000]
    except Exception as e: return f"[Error: {e}]"